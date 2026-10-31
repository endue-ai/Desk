# Builds deck.pptx with python-pptx for the PowerPoint 97-2003 reader: titles, bullets on two levels
# with formatted runs, solid and gradient fills, rotated and flipped shapes, a dashed arrow, a group,
# a cropped PNG and a JPEG, a table, a text box, a slide background, notes and a hidden slide.
# make.sh runs it in a container and converts the deck with LibreOffice. No personal data.
import io
import sys

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
deck = Presentation()
deck.slide_width = Emu(12192000)
deck.slide_height = Emu(6858000)
title_slide, title_content, title_only = deck.slide_layouts[0], deck.slide_layouts[1], deck.slide_layouts[5]
IN = 914400  # EMU per inch; 12700 per point


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def solid(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(rgb)


# 1. Title slide
s = deck.slides.add_slide(title_slide)
s.shapes.title.text = "Quarterly review"
s.placeholders[1].text = "Slides fixture deck"
notes(s, "Welcome. This deck is a test fixture.")

# 2. Bullets on two levels, with bold, italic, underlined and coloured runs
s = deck.slides.add_slide(title_content)
s.shapes.title.text = "Agenda"
body = s.placeholders[1].text_frame
body.text = "Results"
for text, level in [("Revenue grew", 1), ("Costs fell", 1), ("Plans", 0)]:
    p = body.add_paragraph()
    p.text = text
    p.level = level
p = body.add_paragraph()
for text, style in [("Bold", "b"), (" and ", ""), ("italic", "i"), (" and ", ""), ("red", "c"), (" and ", ""), ("under", "u")]:
    r = p.add_run()
    r.text = text
    r.font.bold = True if style == "b" else None
    r.font.italic = True if style == "i" else None
    r.font.underline = True if style == "u" else None
    if style == "c":
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
notes(s, "Keep it short.\nTwo lines of notes.")

# 3. Shapes: fills, a gradient, rotation, a flip, a dashed arrow and a group
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Shapes"
rr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(838200), Emu(1825625), Emu(2286000), Emu(1371600))
rr.text_frame.text = "Inside"
ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(3429000), Emu(1825625), Emu(1828800), Emu(1371600))
solid(ov, "C0392B")
ov.line.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
ov.line.width = Pt(2)
turned = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(5562600), Emu(1825625), Emu(1828800), Emu(914400))
solid(turned, "27AE60")
turned.rotation = 30
turned.text_frame.text = "Turned"
upright = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(7696200), Emu(1825625), Emu(1828800), Emu(914400))
upright.rotation = 90
solid(upright, "8E44AD")
tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Emu(9906000), Emu(1825625), Emu(1371600), Emu(1371600))
tri._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
solid(tri, "F39C12")
gr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(838200), Emu(3657600), Emu(3657600), Emu(1371600))
gr.fill.gradient()
gr.fill.gradient_angle = 0
stops = gr.fill.gradient_stops
stops[0].color.rgb = RGBColor(0x1F, 0x4E, 0x79)
stops[1].color.rgb = RGBColor(0x9D, 0xC3, 0xE6)
gr.text_frame.text = "Gradient"
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(5029200), Emu(4343400), Emu(7315200), Emu(4343400))
ln.line.width = Pt(3)
ln.line.color.rgb = RGBColor(0x33, 0x33, 0x33)
ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
grp = s.shapes.add_group_shape()
a = grp.shapes.add_shape(MSO_SHAPE.DIAMOND, Emu(8001000), Emu(3657600), Emu(914400), Emu(914400))
solid(a, "2980B9")
b = grp.shapes.add_shape(MSO_SHAPE.HEXAGON, Emu(9144000), Emu(3657600), Emu(914400), Emu(914400))
solid(b, "16A085")
notes(s, "Shapes keep their geometry.")

# 4. More fills and lines: three stops at 45 degrees, a radial gradient, a freeform, a dotted outline
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Fills"
three = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(838200), Emu(1828800), Emu(3200400), Emu(1828800))
three.fill.gradient()
three.fill.gradient_angle = 45
gs_lst = three.fill._fill._gradFill.gsLst
stops = three.fill.gradient_stops
stops[0].color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
stops[1].color.rgb = RGBColor(0x34, 0x98, 0xDB)
middle = etree.SubElement(gs_lst, qn("a:gs"))
middle.set("pos", "50000")
etree.SubElement(middle, qn("a:srgbClr")).set("val", "F1C40F")
gs_lst[:] = sorted(gs_lst, key=lambda gs: int(gs.get("pos")))
radial = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(4495800), Emu(1828800), Emu(1828800), Emu(1828800))
radial.fill.gradient()
radial.fill.gradient_stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
radial.fill.gradient_stops[1].color.rgb = RGBColor(0x8E, 0x44, 0xAD)
grad = radial.fill._fill._gradFill
grad.remove(grad.find(qn("a:lin")))
path = etree.SubElement(grad, qn("a:path"))
path.set("path", "circle")
etree.SubElement(path, qn("a:fillToRect")).attrib.update({"l": "50000", "t": "50000", "r": "50000", "b": "50000"})
free = s.shapes.build_freeform(Emu(7010400), Emu(3657600), scale=1.0)
free.add_line_segments([(Emu(8229600), Emu(1828800)), (Emu(9448800), Emu(3657600))], close=True)
star = free.convert_to_shape()
solid(star, "1ABC9C")
dotted = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(838200), Emu(4114800), Emu(3200400), Emu(1371600))
dotted.fill.background()
dotted.line.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
dotted.line.width = Pt(4)
dotted.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
notes(s, "Gradients, a freeform and a dotted outline.")

# 5. A cropped PNG and a JPEG
def picture(fmt, colors):
    img = Image.new("RGB", (160, 100), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for (x, y), color in zip([(0, 0), (80, 0), (0, 50), (80, 50)], colors):
        draw.rectangle([x, y, x + 79, y + 49], fill=color)
    buf = io.BytesIO()
    img.save(buf, fmt)
    buf.seek(0)
    return buf


s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Pictures"
png = picture("PNG", [(231, 76, 60), (46, 204, 113), (52, 152, 219), (241, 196, 15)])
pic = s.shapes.add_picture(png, Emu(838200), Emu(1828800), Emu(4572000), Emu(2857500))
pic.crop_left = 0.1
pic.crop_right = 0.1
jpg = picture("JPEG", [(44, 62, 80), (236, 240, 241), (155, 89, 182), (230, 126, 34)])
s.shapes.add_picture(jpg, Emu(6400800), Emu(1828800), Emu(4572000), Emu(2857500))
notes(s, "The first picture is cropped on both sides.")

# 6. A table in the default style
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "A table"
rows = [["Region", "Q1", "Q2"], ["North", "10", "12"], ["South", "8", "9"]]
table = s.shapes.add_table(3, 3, Emu(1524000), Emu(1828800), Emu(9144000), Emu(1828800)).table
for r, row in enumerate(rows):
    for c, text in enumerate(row):
        table.cell(r, c).text = text
notes(s, "A table with a header row.")

# 7. A text box and a slide background
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Text box"
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor(0xFD, 0xF6, 0xE3)
box = s.shapes.add_textbox(Emu(1524000), Emu(2286000), Emu(9144000), Emu(1371600))
p = box.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Centred in Georgia"
r.font.size = Pt(32)
r.font.name = "Georgia"
r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
notes(s, "A cream background.")

# 8. A hidden slide
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Backup slide"
s._element.set("show", "0")
notes(s, "Hidden in the show.")

# Keep only the three layouts in use: LibreOffice makes a master page of every layout.
for layout in [l for l in deck.slide_layouts if l not in (title_slide, title_content, title_only)]:
    deck.slide_layouts.remove(layout)

deck.save(out)
