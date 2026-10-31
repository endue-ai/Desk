# Builds sample.pptx with python-pptx: titles, bullets, numbering, shapes, a gradient, a line with an
# arrow head, a group, a cropped picture, a table, notes and a hidden slide. make.sh runs it in a
# container and converts the deck with LibreOffice. No personal data.
import copy
import io
import sys

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "sample.pptx"
deck = Presentation()
deck.slide_width = Emu(12192000)
deck.slide_height = Emu(6858000)
title_slide, title_content, title_only = deck.slide_layouts[0], deck.slide_layouts[1], deck.slide_layouts[5]


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# 1. Title slide
s = deck.slides.add_slide(title_slide)
s.shapes.title.text = "Quarterly review"
s.placeholders[1].text = "Slides fixture deck"
notes(s, "Welcome. This deck is a test fixture.")

# 2. Bullets on two levels
s = deck.slides.add_slide(title_content)
s.shapes.title.text = "Agenda"
body = s.placeholders[1].text_frame
body.text = "Results"
for text, level in [("Revenue grew", 1), ("Costs fell", 1), ("Plans", 0)]:
    p = body.add_paragraph()
    p.text = text
    p.level = level
notes(s, "Keep it short.\nTwo lines of notes.")

# 3. Shapes, a gradient, a line with an arrow head and a group
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Shapes"
rr = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(838200), Emu(1825625), Emu(2286000), Emu(1371600))
rr.text_frame.text = "Inside"
ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(3429000), Emu(1825625), Emu(1828800), Emu(1371600))
ov.fill.solid()
ov.fill.fore_color.rgb = RGBColor(0xC0, 0x39, 0x2B)
s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(5562600), Emu(2054225), Emu(1828800), Emu(914400))
s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Emu(7696200), Emu(1825625), Emu(1371600), Emu(1371600))
s.shapes.add_shape(MSO_SHAPE.CHEVRON, Emu(9372600), Emu(1825625), Emu(1371600), Emu(1371600))
gr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(838200), Emu(3657600), Emu(3657600), Emu(1371600))
gr.fill.gradient()
gr.fill.gradient_angle = 0
stops = gr.fill.gradient_stops
stops[0].color.rgb = RGBColor(0x1F, 0x4E, 0x79)
stops[1].color.rgb = RGBColor(0x9D, 0xC3, 0xE6)
gr.text_frame.text = "Gradient"
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(5029200), Emu(4343400), Emu(7315200), Emu(4343400))
ln.line.width = Pt(3)
ln.line.color.rgb = RGBColor(0x33, 0x33, 0x33)
lnxml = ln.line._get_or_add_ln()
tail = etree.SubElement(lnxml, qn("a:tailEnd"))
tail.set("type", "triangle")
grp = s.shapes.add_group_shape()
a = grp.shapes.add_shape(MSO_SHAPE.DIAMOND, Emu(8001000), Emu(3657600), Emu(914400), Emu(914400))
b = grp.shapes.add_shape(MSO_SHAPE.HEXAGON, Emu(9144000), Emu(3657600), Emu(914400), Emu(914400))
notes(s, "Shapes keep their geometry.")

# 4. A cropped picture
img = Image.new("RGB", (160, 100), (255, 255, 255))
draw = ImageDraw.Draw(img)
draw.rectangle([0, 0, 79, 49], fill=(231, 76, 60))
draw.rectangle([80, 0, 159, 49], fill=(46, 204, 113))
draw.rectangle([0, 50, 79, 99], fill=(52, 152, 219))
draw.rectangle([80, 50, 159, 99], fill=(241, 196, 15))
buf = io.BytesIO()
img.save(buf, "PNG")
buf.seek(0)
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "A picture"
pic = s.shapes.add_picture(buf, Emu(3048000), Emu(1828800), Emu(6096000), Emu(3810000))
pic.crop_left = 0.1
pic.crop_right = 0.1
notes(s, "The picture is cropped on both sides.")

# 5. A table in the default style
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "A table"
rows = [["Region", "Q1", "Q2"], ["North", "10", "12"], ["South", "8", "9"]]
table = s.shapes.add_table(3, 3, Emu(1524000), Emu(1828800), Emu(9144000), Emu(1828800)).table
for r, row in enumerate(rows):
    for c, text in enumerate(row):
        table.cell(r, c).text = text
notes(s, "A table with a header row.")

# 6. Numbered steps
s = deck.slides.add_slide(title_content)
s.shapes.title.text = "Numbered steps"
body = s.placeholders[1].text_frame
body.text = "First"
for text in ["Second", "Third"]:
    body.add_paragraph().text = text
for p in body.paragraphs:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "457200")
    pPr.set("indent", "-457200")
    num = etree.SubElement(pPr, qn("a:buAutoNum"))
    num.set("type", "arabicPeriod")

# 7. A hidden slide
s = deck.slides.add_slide(title_only)
s.shapes.title.text = "Backup slide"
s._element.set("show", "0")
notes(s, "Hidden in the show.")

# Keep only the three layouts in use: LibreOffice makes a master page of every layout.
for layout in [l for l in deck.slide_layouts if l not in (title_slide, title_content, title_only)]:
    deck.slide_layouts.remove(layout)

deck.save(out)
